from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook, load_workbook

from app.config import Config
from app.models import Inventory, Visitor

VISITORS_HEADERS = [
    "issue_id",
    "full_name",
    "mobile",
    "stay_days",
    "gadda_qty",
    "rajai_qty",
    "deposit_amount",
    "issue_datetime",
    "return_datetime",
    "status",
]

INVENTORY_HEADERS = [
    "total_gadda",
    "available_gadda",
    "total_rajai",
    "available_rajai",
    "deposit_per_item",
]


def _ensure_file(path: Path, creator):
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        creator(path)


def _create_visitors_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Visitors"
    ws.append(VISITORS_HEADERS)
    wb.save(path)


def _create_inventory_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(INVENTORY_HEADERS)
    ws.append([
        Config.DEFAULT_TOTAL_GADDA,
        Config.DEFAULT_TOTAL_GADDA,
        Config.DEFAULT_TOTAL_RAJAI,
        Config.DEFAULT_TOTAL_RAJAI,
        Config.DEFAULT_DEPOSIT_PER_ITEM,
    ])
    wb.save(path)


def _parse_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        try:
            return datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None


def ensure_storage():
    _ensure_file(Config.VISITORS_EXCEL, _create_visitors_file)
    _ensure_file(Config.INVENTORY_EXCEL, _create_inventory_file)


def load_visitors() -> List[Visitor]:
    ensure_storage()
    wb = load_workbook(Config.VISITORS_EXCEL)
    ws = wb["Visitors"]
    visitors: List[Visitor] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        visitors.append(_visitor_from_row(row))
    return visitors


def _visitor_from_row(row):
    return Visitor(
        issue_id=str(row[0]),
        full_name=str(row[1] or ""),
        mobile=str(row[2] or ""),
        stay_days=int(row[3] or 0),
        gadda_qty=int(row[4] or 0),
        rajai_qty=int(row[5] or 0),
        deposit_amount=int(row[6] or 0),
        issue_datetime=_parse_datetime(row[7]) or datetime.now(timezone.utc),
        return_datetime=_parse_datetime(row[8]),
        status=str(row[9] or "active"),
    )


def load_inventory() -> Inventory:
    ensure_storage()
    wb = load_workbook(Config.INVENTORY_EXCEL)
    ws = wb["Inventory"]
    row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
    if row is None or row[0] is None:
        inventory = Inventory(
            total_gadda=Config.DEFAULT_TOTAL_GADDA,
            available_gadda=Config.DEFAULT_TOTAL_GADDA,
            total_rajai=Config.DEFAULT_TOTAL_RAJAI,
            available_rajai=Config.DEFAULT_TOTAL_RAJAI,
            deposit_per_item=Config.DEFAULT_DEPOSIT_PER_ITEM,
        )
        save_inventory(inventory)
        return inventory
    return Inventory(
        total_gadda=int(row[0] or 0),
        available_gadda=int(row[1] or 0),
        total_rajai=int(row[2] or 0),
        available_rajai=int(row[3] or 0),
        deposit_per_item=int(row[4] or Config.DEFAULT_DEPOSIT_PER_ITEM),
    )


def save_inventory(inventory: Inventory):
    ensure_storage()
    wb = load_workbook(Config.INVENTORY_EXCEL)
    ws = wb["Inventory"]
    if ws.max_row < 2:
        ws.append([
            inventory.total_gadda,
            inventory.available_gadda,
            inventory.total_rajai,
            inventory.available_rajai,
            inventory.deposit_per_item,
        ])
    else:
        row = ws[2]
        row[0].value = inventory.total_gadda
        row[1].value = inventory.available_gadda
        row[2].value = inventory.total_rajai
        row[3].value = inventory.available_rajai
        row[4].value = inventory.deposit_per_item
    wb.save(Config.INVENTORY_EXCEL)


def append_visitor(visitor: Visitor):
    ensure_storage()
    wb = load_workbook(Config.VISITORS_EXCEL)
    ws = wb["Visitors"]
    ws.append([
        visitor.issue_id,
        visitor.full_name,
        visitor.mobile,
        visitor.stay_days,
        visitor.gadda_qty,
        visitor.rajai_qty,
        visitor.deposit_amount,
        visitor.issue_datetime.isoformat(),
        visitor.return_datetime.isoformat() if visitor.return_datetime else "",
        visitor.status,
    ])
    wb.save(Config.VISITORS_EXCEL)


def update_visitor(issue_id: str, status: Optional[str] = None, return_datetime: Optional[datetime] = None):
    ensure_storage()
    wb = load_workbook(Config.VISITORS_EXCEL)
    ws = wb["Visitors"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == issue_id:
            if status is not None:
                row[9].value = status
            if return_datetime is not None:
                row[8].value = return_datetime.isoformat()
            wb.save(Config.VISITORS_EXCEL)
            return _visitor_from_row([cell.value for cell in row])
    raise ValueError(f"Visitor not found: {issue_id}")


def find_visitor_by_issue_or_mobile(query: str, active_only: bool = False) -> Optional[Visitor]:
    q = query.strip().lower()
    for visitor in reversed(load_visitors()):
        if active_only and visitor.status != "active":
            continue
        if visitor.issue_id.lower() == q or visitor.mobile.lower() == q:
            return visitor
    return None


def get_next_issue_number() -> int:
    visitors = load_visitors()
    highest = 0
    for visitor in visitors:
        try:
            parts = visitor.issue_id.split("-")
            if parts:
                num = int(parts[-1])
                highest = max(highest, num)
        except ValueError:
            continue
    return highest + 1


def active_issued_counts() -> tuple[int, int]:
    visitors = load_visitors()
    g = sum(v.gadda_qty for v in visitors if v.status == "active")
    r = sum(v.rajai_qty for v in visitors if v.status == "active")
    return g, r


def all_visitors_sorted(desc: bool = True) -> List[Visitor]:
    visitors = load_visitors()
    visitors.sort(key=lambda item: item.issue_datetime or datetime.min.replace(tzinfo=timezone.utc), reverse=desc)
    return visitors


def visitor_summary():
    visitors = load_visitors()
    total = len(visitors)
    active = sum(1 for v in visitors if v.status == "active")
    returned = sum(1 for v in visitors if v.status == "returned")
    total_gadda = sum(v.gadda_qty for v in visitors)
    total_rajai = sum(v.rajai_qty for v in visitors)
    total_deposit = sum(v.deposit_amount for v in visitors)
    return {
        "total": total,
        "active": active,
        "returned": returned,
        "total_gadda": total_gadda,
        "total_rajai": total_rajai,
        "total_deposit": total_deposit,
    }
